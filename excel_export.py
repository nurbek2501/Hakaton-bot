# -*- coding: utf-8 -*-
"""
Barcha anketalarni bitta chiroyli Excel (.xlsx) faylga yig‘ish (/export buyrug‘i uchun).

1-varaq «Ro‘yxat»     - har bir ishtirokchi bitta qatorda, filtr va qotirilgan sarlavha bilan
2-varaq «Statistika»  - soha, ishtirok shakli, faoliyat turi va h.k. bo‘yicha sonlar
"""
import io
import math
from collections import Counter
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from questions import OTHER, QUESTIONS, is_applicable
from storage import TZ

TITLE = "«Xatirchi raqamli yoshlari xakatoni – 2026» — ishtirokchilar ro‘yxati"

# Ustun nomlari (qisqa). Yangi savol qo‘shilsa va bu yerda bo‘lmasa - savol matni ishlatiladi
COLUMN_NAMES = {
    "fio": "F.I.Sh.", "birth": "Tug‘ilgan sana", "gender": "Jinsi",
    "citizenship": "Fuqaroligi", "residence": "Xatirchida ro‘yxatda",
    "address": "Yashash manzili", "phone": "Telefon", "username": "Telegram",
    "email": "E-mail", "activity": "Faoliyat turi", "workplace": "O‘qish / ish joyi",
    "participation": "Ishtirok shakli", "team_name": "Jamoa nomi",
    "team_members": "Jamoa a’zolari", "project_name": "Loyiha nomi", "sphere": "Soha",
    "problem": "Muammo", "solution": "Yechim", "readiness": "Tayyorlik darajasi",
    "tech": "Texnologiyalar", "link": "Havolalar va fayllar",
    "agreement": "Talablarga rozilik", "source": "Qayerdan bildi",
}
# Ustun kengligi savol turiga qarab
WIDTHS = {"fio": 28, "date": 13, "choice": 24, "text": 28, "phone": 18, "username": 18,
          "email": 26, "members": 45, "longtext": 50, "link_or_file": 45, "confirm": 12}

# Ranglar
C_HEADER = "1F4E78"
C_ZEBRA = "F2F6FA"
C_FAILED = "FDE2E1"
C_BORDER = "D0D7DE"

THIN = Side(style="thin", color=C_BORDER)
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor=C_HEADER)
TOP_WRAP = Alignment(vertical="top", wrap_text=True)
TOP_CENTER = Alignment(vertical="top", horizontal="center", wrap_text=True)


def _clean(v):
    """Excel qabul qilmaydigan boshqaruv belgilarini olib tashlaydi."""
    return ILLEGAL_CHARACTERS_RE.sub("", v) if isinstance(v, str) else v


def _set(cell, value):
    """Qiymat yozadi. "=" bilan boshlangan matn formula bo‘lib ishlab ketmasligi uchun
    doim oddiy matn sifatida saqlanadi (Excel formula injection’dan himoya)."""
    cell.value = _clean(value)
    if isinstance(value, str) and value.startswith("="):
        cell.data_type = "s"


def _links(value):
    """Loyiha havola/fayllari: (matn, birinchi_havola)."""
    items = [{"t": "url", "v": value}] if isinstance(value, str) else (value or [])
    lines, first = [], None
    for it in items:
        if it.get("t") == "file":
            link = it.get("link")
            lines.append(f"{it['name']}: {link}" if link else f"{it['name']} (Telegram, adminlarda)")
        else:
            link = it.get("v")
            lines.append(link)
        first = first or link
    return "\n".join(lines), first


def _cell_value(q, answers):
    """Savol javobini Excel katagi uchun tayyorlaydi: (qiymat, havola)."""
    if not is_applicable(q, answers):
        return None, None
    v = answers.get(q["key"])
    if v in (None, "", []):
        return None, None
    t = q["type"]
    if t == "date":
        d, m, y = map(int, v.split("."))
        return date(y, m, d), None
    if t == "confirm":
        return "Ha", None
    if t == "link_or_file":
        return _links(v)
    if v == OTHER and q.get("other"):
        return f"{OTHER}: {answers.get(q['key'] + '__other', '')}", None
    return v, None


def _age(born, today):
    return today.year - born.year - ((today.month, today.day) < (born.month, born.day))


def _lines_needed(value, width):
    if value is None:
        return 1
    text = value if isinstance(value, str) else str(value)
    return sum(max(1, math.ceil(len(part) / max(width - 2, 1))) for part in text.split("\n"))


def _columns():
    """[(sarlavha, kenglik, savol yoki maxsus kalit)]"""
    cols = [("№", 5, "#")]
    for q in QUESTIONS:
        name = COLUMN_NAMES.get(q["key"]) or q["text"].split(" (")[0]
        cols.append((name, WIDTHS.get(q["type"], 24), q))
        if q["type"] == "date":
            cols.append(("Yoshi", 7, "age"))
    cols += [("Telegram ID", 14, "uid"), ("Ro‘yxatdan o‘tgan vaqt", 18, "ts"), ("Holat", 22, "status")]
    return cols


def _sheet_list(ws, rows, now):
    cols = _columns()
    last_col = get_column_letter(len(cols))
    header_row = 4

    # Sarlavha
    ws["A1"] = TITLE
    ws["A1"].font = Font(bold=True, size=14, color=C_HEADER)
    ws.merge_cells(f"A1:{last_col}1")
    ws["A2"] = f"Yuklab olingan: {now:%d.%m.%Y %H:%M}   ·   Jami: {len(rows)} ta anketa"
    ws["A2"].font = Font(italic=True, color="666666")
    ws.merge_cells(f"A2:{last_col}2")

    for c, (name, width, _) in enumerate(cols, 1):
        cell = ws.cell(row=header_row, column=c, value=name)
        cell.font, cell.fill, cell.border = HEADER_FONT, HEADER_FILL, BORDER
        cell.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = width
    ws.row_dimensions[header_row].height = 32

    today = now.date()
    for n, (uid, ts, status, answers) in enumerate(rows, 1):
        r = header_row + n
        failed = not status.startswith("Formaga")
        fill = PatternFill("solid", fgColor=C_FAILED if failed else C_ZEBRA) if (failed or n % 2 == 0) else None
        birth = None
        lines = 1
        for c, (_, width, src) in enumerate(cols, 1):
            cell = ws.cell(row=r, column=c)
            link, align = None, TOP_WRAP
            if src == "#":
                value, align = n, TOP_CENTER
            elif src == "age":
                value, align = (_age(birth, today) if birth else None), TOP_CENTER
            elif src == "uid":
                value = str(uid)
            elif src == "ts":
                try:
                    value = datetime.strptime(ts, "%Y-%m-%d %H:%M:%S")
                    cell.number_format = "DD.MM.YYYY HH:MM"
                except (TypeError, ValueError):
                    value = ts
            elif src == "status":
                value = ("⚠️ " if failed else "✅ ") + status
            else:
                value, link = _cell_value(src, answers)
                if src["type"] == "date" and value:
                    birth = value
                    cell.number_format = "DD.MM.YYYY"
                    align = TOP_CENTER
            _set(cell, value)
            if link:
                cell.hyperlink = link
                cell.font = Font(color="0563C1", underline="single")
            cell.alignment, cell.border = align, BORDER
            if fill:
                cell.fill = fill
            lines = max(lines, min(6, _lines_needed(value, width)))
        ws.row_dimensions[r].height = 15 * lines + 3

    last_row = header_row + max(len(rows), 1)
    ws.auto_filter.ref = f"A{header_row}:{last_col}{last_row}"
    ws.freeze_panes = f"C{header_row + 1}"  # sarlavha va №, F.I.Sh. ustunlari qotiriladi


def _sheet_stats(ws, rows):
    ws["A1"] = "Statistika"
    ws["A1"].font = Font(bold=True, size=14, color=C_HEADER)
    ws["A2"] = f"Jami anketalar: {len(rows)} ta"
    ws["A2"].font = Font(bold=True)
    ws.column_dimensions["A"].width = 55
    ws.column_dimensions["B"].width = 10
    ws.column_dimensions["C"].width = 10

    r = 4
    for q in QUESTIONS:
        if q["type"] != "choice":
            continue
        counts = Counter()
        for _, _, _, answers in rows:
            if is_applicable(q, answers):
                counts[answers.get(q["key"]) or "Ko‘rsatilmagan"] += 1
        total = sum(counts.values())
        if not total:
            continue
        for c, name in enumerate((COLUMN_NAMES.get(q["key"], q["text"]), "Soni", "Ulushi"), 1):
            cell = ws.cell(row=r, column=c, value=name)
            cell.font, cell.fill, cell.border = HEADER_FONT, HEADER_FILL, BORDER
        r += 1
        # Variantlar formadagi tartibda, keyin «Boshqa» va ko‘rsatilmaganlar
        order = [v for _, v in q["options"]] + [OTHER, "Ko‘rsatilmagan"]
        for value in sorted(counts, key=lambda v: order.index(v) if v in order else len(order)):
            cells = (ws.cell(row=r, column=1, value=_clean(value)),
                     ws.cell(row=r, column=2, value=counts[value]),
                     ws.cell(row=r, column=3, value=counts[value] / total))
            cells[2].number_format = "0%"
            for cell in cells:
                cell.border = BORDER
            r += 1
        r += 1


def build_xlsx(rows):
    """rows: storage.export_rows() natijasi. Qaytaradi: .xlsx fayl baytlari."""
    now = datetime.now(TZ).replace(tzinfo=None)
    wb = Workbook()
    ws = wb.active
    ws.title = "Ro‘yxat"
    _sheet_list(ws, rows, now)
    _sheet_stats(wb.create_sheet("Statistika"), rows)
    wb.properties.title = TITLE
    wb.properties.creator = "Xakaton bot"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
